from functools import partial

from bokeh.io import show
from bokeh.io import output_notebook
from bokeh.layouts import column, gridplot, row
from bokeh.models import CustomJS, Select, Button, ColumnDataSource, Div
from bokeh.models.widgets import HTMLTemplateFormatter, TableColumn
from bokeh.models.widgets.tables import DataTable
from bokeh.themes import Theme
from bokeh.plotting import curdoc
from openpyxl import load_workbook
import pandas as pd

from IPython.display import display

from team import Team
import espn_api as api

# user selection later?
season = '2023'
league_id = 32274580
scoring_period = '2023'

# get list of all team names for display
team_names = api.get_team_names(season, league_id)
team_info = api.get_team_info(season, league_id)
team_dict = api.create_team_dictionary(team_info[1])

# team objects that hold relevant team info
# initialized to first two teams
team1 = Team(season, league_id, scoring_period)
team2 = Team(season, league_id, scoring_period)
team1.set_name(team_names[0])
team1.set_team_id(team_dict[team_names[0]])
team2.set_name(team_names[1])
team2.set_team_id(team_dict[team_names[1]])

def store_team(attr, old, new, team):
    team.set_name(new)
    team.set_team_id(team_dict[new])
    return

team1_select = Select(title='Team 1', options=team_names, value=team_names[0])
team1_select.on_change('value', partial(store_team, team=team1))


team2_select = Select(title='Team 2', options=team_names, value=team_names[1])
team2_select.on_change('value', partial(store_team, team=team2))

def store_scoring_period(attr, old, new, tm1, tm2):
    tm1.set_scoring_period(new)
    tm2.set_scoring_period(new)
    return

project_select = Select(title='Projection Type', options=['2023', 'Last 7 Days', 'Last 15 Days', 'Last 30 Days', '2023 Projections'], value='2023')
project_select.on_change('value', partial(store_scoring_period, tm1=team1, tm2=team2))

# curdoc().add_root(team1_select)
# curdoc().add_root(team2_select)
# curdoc().add_root(project_select)
options_row = row(team1_select, team2_select, project_select)
curdoc().add_root(options_row)

# on change + default value not working

gen_proj_button = Button(label='Generate Projections', button_type='success')

df_cols = ['Player', 'FGM', 'FGA', 'FG%', 'FTM', 'FTA', 'FT%', '3PM', 'REB', 'AST', 'STL', 'BLK', 'TO', 'PTS']
dt_cols = [TableColumn(field=val, title=val, width=150) if val == 'Player'
           else TableColumn(field=val, title=val, width=40) for val in df_cols]

def get_empty_cds():
    empty_df = pd.DataFrame([[''] * 14], columns=df_cols)
    return ColumnDataSource(dict(empty_df))


team1_table = DataTable(source=get_empty_cds(), columns=dt_cols, reorderable=False, editable=True, index_position=None,
                        fit_columns=False, sortable=False, width=680, height=500)
team2_table = DataTable(source=get_empty_cds(), columns=dt_cols, reorderable=False, editable=True, index_position=None,
                        fit_columns=False, sortable=False, width=680, height=500)

team1_div = Div(text='Team 1')
team2_div = Div(text='Team 2')
diff_div = Div(text='Differences')

def generate_projections(team1, team1_div, team2, team2_div):
    team1.get_team_data()
    team1_table.source.data = team1.output_team_data()
    team1_div.text = team1.get_name()

    team2.get_team_data()
    team2_table.source.data = team2.output_team_data()
    team2_div.text = team2.get_name()

    diff_table.source.data = team1.get_stat_diff(team2)
    return

gen_proj_button.on_click(
    partial(generate_projections, team1=team1, team1_div=team1_div, team2=team2, team2_div=team2_div))

out_proj_button = Button(label='Output Team 1 Comparisons ', button_type='success')

def format_excel():
    wb = load_workbook('diffs.xlsx')
    ws = wb.active()
    return

def output_projections(team1, team_dict):
    team1.get_team_data()
    other_team = Team(season, league_id, team1.get_scoring_period())
    team_diffs = []
    for team in team_names:
        if team == team1.get_name():
            continue
        other_team.set_name(team)
        other_team.set_team_id(team_dict[team])
        other_team.get_team_data()
        other_team_diff = team1.get_stat_diff(other_team).iloc[[2]].copy()
        other_team_diff.insert(0, 'Team', team)
        other_team_diff.drop(['Player', 'FGM', 'FGA', 'FTM', 'FTA'], axis=1, inplace=True)
        team_diffs.append(other_team_diff)
    df_diffs = pd.concat(team_diffs)
    df_diffs.to_excel('diffs.xlsx', index=False)
    print('Team 1 Comparisons Generated.')
    return

out_proj_button.on_click(
    partial(output_projections, team1=team1, team_dict=team_dict))

curdoc().add_root(row(gen_proj_button, out_proj_button))

team_tables = row(column(team1_div, team1_table), column(team2_div, team2_table))
curdoc().add_root(team_tables)

template="""                
            <p style="color:<%= 
                (function colorfromint(){
                    if (1 < Math.abs(cola - colb) && Math.abs(cola - colb) < 10)
                        {return('green')}
                    else if (10 < Math.abs(cola - colb) && Math.abs(cola - colb) < 40)
                        {return('blue')}
                    else 
                        {return('red')}
                    }()) %>;"> 
                <%= value %>
            </p>
            """
formatter =  HTMLTemplateFormatter(template=template)

# add css classes?
diff_table = DataTable(source=get_empty_cds(), columns=dt_cols, reorderable=False, editable=False, sortable=False,
                       selectable=True, index_position=None, fit_columns=False, width=680)
# diff_table = DataTable(source=get_empty_cds(), columns=dt_cols, editable=True,
#                        index_position=None, fit_columns=False, width=680)
curdoc().add_root(diff_table)





#     doc.theme = Theme(json=yaml.load("""
#         attrs:
#             Figure:
#                 background_fill_color: "#DDDDDD"
#                 outline_line_color: white
#                 toolbar_location: above
#                 height: 500
#                 width: 800
#             Grid:
#                 grid_line_dash: [6, 4]
#                 grid_line_color: white
#     """, Loader=yaml.FullLoader))